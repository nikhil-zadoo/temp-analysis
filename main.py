from collections import Counter
from docx import Document
import docx2txt

plain_text = docx2txt.process("kashmiri.docx")
list_of_words = plain_text.split()
#print(Counter(list_of_words))
counter_list_of_words = Counter(list_of_words)
elements = counter_list_of_words.items()
# for a, b in elements:
#     print(a)
#     print(b)

doc = Document()
# Create and Name Table Heading
table = doc.add_table(rows=1, cols=2)
cell1 = table.cell(0, 0)
cell1.text = 'Word'
cell2 = table.cell(0, 1)
cell2.text = 'Frequency'

#Iterate over collection elements and append to table craeted
for word, frequency in elements:
    cell = table.add_row().cells
    cell[0].text = str(word)
    cell[1].text = str(frequency)
doc.save("results.docx")